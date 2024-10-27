import os
import pandas as pd
import numpy as np
from sklearn.kernel_ridge import KernelRidge
from sklearn.metrics import mean_squared_error
from sklearn.preprocessing import StandardScaler
from openpyxl import Workbook
from openpyxl.styles import PatternFill

def kernel_ridge_regression(train_data, test_data):
    X_train = train_data[['Qv', 'DP', 'RPM']]
    y_train = train_data['M1']
    
    X_test = test_data[['Qv', 'DP', 'RPM']]
    y_test = test_data['M1']

    # 标准化特征
    scaler = StandardScaler()
    X_train_scaled = scaler.fit_transform(X_train)
    X_test_scaled = scaler.transform(X_test)

    # 创建并训练内核岭回归模型
    model = KernelRidge(alpha=1.0, kernel='rbf')
    model.fit(X_train_scaled, y_train)

    # 预测并计算指标
    y_pred = model.predict(X_test_scaled)
    mse = mean_squared_error(y_test, y_pred)

    return mse

def main():
    # 设置文件路径
    current_dir = os.path.dirname(os.path.abspath(__file__))
    file_path = os.path.join(current_dir, '..', 'data_base.xlsx')
    output_path = os.path.join(current_dir, 'RidgeRegression.xlsx')

    # 定义sheet名称
    sheet_names = ['CVAF', 'CVAR', 'HFF', 'HDF']

    # 读取所有数据
    all_data = pd.DataFrame()
    for sheet_name in sheet_names:
        df = pd.read_excel(file_path, sheet_name=sheet_name)
        df['Sheet'] = sheet_name
        all_data = pd.concat([all_data, df])

    # 创建一个新的Excel工作簿
    wb = Workbook()
    wb.remove(wb.active)  # 删除默认创建的sheet

    # 遍历每个sheet
    for sheet_name in sheet_names:
        sheet_data = all_data[all_data['Sheet'] == sheet_name].copy()
        ws = wb.create_sheet(sheet_name)

        # 添加表头
        headers = list(sheet_data.columns) + ['MSE']
        ws.append(headers)

        # 遍历每一行数据作为测试集
        for idx, row in sheet_data.iterrows():
            test_data = pd.DataFrame([row])
            train_data = all_data[all_data.index != idx]

            # 执行内核岭回归分析
            mse = kernel_ridge_regression(train_data, test_data)

            # 将结果添加到工作表中
            result_row = list(row) + [mse]
            ws.append(result_row)

            # 如果MSE大于16，将整行标记为红色
            if mse > 16:
                for cell in ws[ws.max_row]:
                    cell.fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")

    # 保存结果
    wb.save(output_path)
    print(f"结果已保存到 {output_path}")

if __name__ == "__main__":
    main()
